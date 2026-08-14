import asyncio
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

import pypdfium2 as pdfium
import pytesseract
from docx import Document as DocxDocument
from PIL import Image
from pypdf import PdfReader

from .models import ExtractionMethod


class OcrError(RuntimeError):
    """Raised when a local document cannot be converted into page text."""


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int
    text: str
    extraction_method: ExtractionMethod


class PageTextExtractor(Protocol):
    """Extract ordered page text behind a replaceable OCR provider boundary."""

    async def extract(self, path: Path) -> list[ExtractedPage]: ...


class LocalPageTextExtractor:
    """Extract native text first and use local Tesseract OCR when needed."""

    async def extract(self, path: Path) -> list[ExtractedPage]:
        return await asyncio.to_thread(self._extract_sync, path)

    def _extract_sync(self, path: Path) -> list[ExtractedPage]:
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                return self._extract_pdf(path)
            if suffix in {".png", ".jpg", ".jpeg"}:
                return self._extract_image(path)
            if suffix == ".docx":
                return self._extract_docx(path)
        except OcrError:
            raise
        except Exception as exc:
            raise OcrError("document_text_extraction_failed") from exc
        raise OcrError("unsupported_document_type")

    def _extract_pdf(self, path: Path) -> list[ExtractedPage]:
        reader = PdfReader(path)
        if not reader.pages:
            raise OcrError("document_has_no_pages")

        native_text = [(page.extract_text() or "").strip() for page in reader.pages]
        rendered_pdf = None
        pages: list[ExtractedPage] = []
        try:
            for index, text in enumerate(native_text):
                if text:
                    pages.append(
                        ExtractedPage(index + 1, text, ExtractionMethod.native)
                    )
                    continue

                if rendered_pdf is None:
                    rendered_pdf = pdfium.PdfDocument(str(path))
                pdf_page = rendered_pdf[index]
                bitmap = pdf_page.render(scale=2)
                image = bitmap.to_pil()
                try:
                    ocr_text = self._ocr_image(image)
                finally:
                    image.close()
                    bitmap.close()
                    pdf_page.close()
                pages.append(
                    ExtractedPage(index + 1, ocr_text, ExtractionMethod.ocr)
                )
        finally:
            if rendered_pdf is not None:
                rendered_pdf.close()
        return pages

    def _extract_image(self, path: Path) -> list[ExtractedPage]:
        with Image.open(path) as source:
            frame_count = getattr(source, "n_frames", 1)
            pages = []
            for index in range(frame_count):
                source.seek(index)
                frame = source.convert("RGB")
                try:
                    text = self._ocr_image(frame)
                finally:
                    frame.close()
                pages.append(ExtractedPage(index + 1, text, ExtractionMethod.ocr))
        return pages

    def _extract_docx(self, path: Path) -> list[ExtractedPage]:
        document = DocxDocument(str(path))
        text = "\n".join(
            paragraph.text for paragraph in document.paragraphs if paragraph.text
        ).strip()
        return [ExtractedPage(1, text, ExtractionMethod.native)]

    @staticmethod
    def _ocr_image(image: Image.Image) -> str:
        return pytesseract.image_to_string(image).strip()
